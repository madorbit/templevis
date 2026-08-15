"""Tests for veil (V code) extraction and staffing sheet generation."""

import os

import pytest
from openpyxl import load_workbook

from templevis.staffing_generator import StaffingAssignmentsGenerator, _display_names
from templevis.veil_extractor import VeilPDFExtractor, parse_time_of_day

TEST_DATA_DIR = 'memory-bank'
BROTHERS_PDF = os.path.join(TEST_DATA_DIR, 'Brothers aug 5.pdf')
SISTERS_PDF = os.path.join(TEST_DATA_DIR, 'Sisters Aug 5.pdf')
TEMPLATE = os.path.join(TEST_DATA_DIR, 'StaffingAssignments.xlsx')
SESSION_TIMES = ['6:05', '7:20', '8:05', '8:50']


def _require(path):
    if not os.path.exists(path):
        pytest.skip(f"Test file not found: {path}")
    return path


@pytest.fixture(scope="module")
def brothers():
    extractor = VeilPDFExtractor(_require(BROTHERS_PDF))
    extractor.extract()
    return extractor


@pytest.fixture(scope="module")
def sisters():
    extractor = VeilPDFExtractor(_require(SISTERS_PDF))
    extractor.extract()
    return extractor


def test_parse_time_of_day_assumes_evening():
    assert parse_time_of_day('6:05 Veil') == 18 * 60 + 5
    assert parse_time_of_day('nope') is None


def test_veil_boxes_span_forty_five_minutes(brothers):
    veil = [a for a in brothers.assignments if a.kind == 'V']
    assert veil
    assert all(a.end_minutes - a.start_minutes == 45 for a in veil)


def test_each_session_has_presenters_and_one_officiator(brothers):
    for time_text in SESSION_TIMES:
        session = brothers.session_at(parse_time_of_day(time_text))
        assert session is not None, time_text
        assert len(session['presenters']) == 14
        assert session['officiator']


def test_officiator_matches_session_slot(sisters):
    session = sisters.session_at(parse_time_of_day('8:05'))
    assert session['code'] == 'V-1'
    assert session['officiator'] == 'Schoenfeld, Marsha'


def test_non_numeric_veil_codes_are_not_presenters(sisters):
    session = sisters.session_at(parse_time_of_day('6:05'))
    assert 'Gillespie, LeAnn' not in session['presenters']


def test_trainees_are_not_presenters(sisters):
    session = sisters.session_at(parse_time_of_day('8:50'))
    assert session['presenters'] == [
        'Durkin, Inger', 'Hunsaker, Debra', 'Taylor, Lisa', 'Winters, Kamla'
    ]


def test_display_names_drop_first_name_unless_ambiguous():
    assert _display_names(['Jones, Al', 'Smith, Kevin']) == ['Jones', 'Smith']
    assert _display_names(['Smith, Kevin', 'Smith, Sam']) == ['Smith, K.', 'Smith, S.']


def test_generate_fills_every_section(tmp_path):
    _require(BROTHERS_PDF)
    _require(SISTERS_PDF)
    _require(TEMPLATE)

    output = str(tmp_path / 'staffing.xlsx')
    generator = StaffingAssignmentsGenerator(TEMPLATE)
    generator.generate(BROTHERS_PDF, SISTERS_PDF, output_path=output)

    sheet = load_workbook(output).active
    sections = generator._find_sections(sheet)

    assert len(sections) == 4
    for section in sections:
        brothers_written = [
            sheet.cell(row=row, column=section.brother_col).value
            for row in section.upper_rows
        ] + [
            sheet.cell(row=row, column=col).value
            for col in (section.brother_col, section.sister_col)
            for row in section.lower_rows
        ]
        sisters_written = [
            sheet.cell(row=row, column=section.sister_col).value
            for row in section.upper_rows
        ]
        assert len([n for n in brothers_written if n]) == 14
        assert len([n for n in sisters_written if n]) >= 4

        for col in (section.brother_col, section.sister_col):
            eo_value = sheet.cell(row=section.eo_row, column=col).value
            assert eo_value.startswith('EO:') and len(eo_value) > len('EO:')
