"""
Staffing Assignments Generator

Fills the StaffingAssignments.xlsx template from the brother and sister
schedule PDFs, using the veil (V) assignment boxes to decide who is serving at
each printed veil session time.
"""

import os
import re
from collections import Counter

from openpyxl import load_workbook

from .veil_extractor import VeilPDFExtractor, parse_time_of_day

SESSION_PATTERN = re.compile(r'^\s*(\d{1,2}:\d{2})\s+Veil\s*$', re.IGNORECASE)

# Above the 'EO:' row the two columns are brothers and sisters; below it both
# columns hold brothers.
TYPICAL_UPPER_USED = 4          # Typical week leaves the fifth upper row blank
TYPICAL_LOWER_PER_COLUMN = 5    # Typical week leaves the last lower row blank
TYPICAL_TOTAL = 14              # Brothers assigned to a veil session in a typical week

DEFAULT_TEMPLATE = 'StaffingAssignments.xlsx'
TEMPLATE_SEARCH_DIRS = ('.', 'memory-bank')


class StaffingSection:
    """One veil session block within the staffing sheet."""

    def __init__(self, time_row, header_row, eo_row, notes_row, brother_col, session_time):
        self.time_row = time_row
        self.header_row = header_row
        self.eo_row = eo_row
        self.notes_row = notes_row
        self.brother_col = brother_col
        self.sister_col = brother_col + 1
        self.session_time = session_time

    @property
    def upper_rows(self):
        return list(range(self.header_row + 1, self.eo_row))

    @property
    def lower_rows(self):
        return list(range(self.eo_row + 1, self.notes_row))


class StaffingAssignmentsGenerator:
    """Generates a filled copy of the veil staffing assignments workbook."""

    def __init__(self, template_path=DEFAULT_TEMPLATE):
        self.template_path = _resolve_template(template_path)

    def generate(self, brothers_pdf=None, sisters_pdf=None, output_path=None,
                 output_dir='output'):
        """Fill the template from one or both schedule PDFs.

        Args:
            brothers_pdf: Path to the brother volunteer schedule PDF.
            sisters_pdf: Path to the sister volunteer schedule PDF.
            output_path: Explicit output file path.
            output_dir: Directory used when output_path is not given.

        Returns:
            str: path of the workbook that was written.
        """
        if not brothers_pdf and not sisters_pdf:
            raise ValueError("At least one schedule PDF is required")

        brothers = self._load(brothers_pdf)
        sisters = self._load(sisters_pdf)

        workbook = load_workbook(self.template_path)
        sheet = workbook.active

        schedule_date = None
        for source in (brothers, sisters):
            if source is not None and source.schedule_date:
                schedule_date = source.schedule_date
                break
        if schedule_date:
            self._set_date(sheet, schedule_date)

        for section in self._find_sections(sheet):
            self._fill_section(sheet, section, brothers, sisters)

        if output_path is None:
            suffix = schedule_date.isoformat() if schedule_date else 'output'
            os.makedirs(output_dir, exist_ok=True)
            output_path = os.path.join(
                output_dir, f"StaffingAssignments_{suffix}.xlsx"
            )
        else:
            parent = os.path.dirname(output_path)
            if parent:
                os.makedirs(parent, exist_ok=True)

        workbook.save(output_path)
        return output_path

    def _load(self, pdf_path):
        if not pdf_path:
            return None
        extractor = VeilPDFExtractor(pdf_path)
        extractor.extract()
        return extractor

    def _find_sections(self, sheet):
        sections = []
        for row in range(1, sheet.max_row + 1):
            for col in range(1, sheet.max_column + 1):
                value = sheet.cell(row=row, column=col).value
                if not isinstance(value, str):
                    continue
                match = SESSION_PATTERN.match(value)
                if not match:
                    continue

                header_row = row + 1
                eo_row = self._find_label_row(sheet, col, header_row + 1, 'EO')
                if eo_row is None:
                    continue
                notes_row = self._find_label_row(sheet, col, eo_row + 1, 'Notes')
                if notes_row is None:
                    continue

                sections.append(StaffingSection(
                    time_row=row,
                    header_row=header_row,
                    eo_row=eo_row,
                    notes_row=notes_row,
                    brother_col=col,
                    session_time=parse_time_of_day(match.group(1)),
                ))
        return sections

    def _find_label_row(self, sheet, col, start_row, label):
        for row in range(start_row, sheet.max_row + 1):
            value = sheet.cell(row=row, column=col).value
            if isinstance(value, str) and value.strip().lower().startswith(label.lower()):
                return row
        return None

    def _set_date(self, sheet, schedule_date):
        for row in range(1, sheet.max_row + 1):
            for col in range(1, sheet.max_column + 1):
                value = sheet.cell(row=row, column=col).value
                if isinstance(value, str) and value.strip().lower() == 'date':
                    sheet.cell(row=row, column=col + 1).value = schedule_date
                    return

    def _fill_section(self, sheet, section, brothers, sisters):
        overflow = []

        brother_session = self._session(brothers, section.session_time)
        if brother_session:
            overflow += self._write_brothers(sheet, section, brother_session)

        sister_session = self._session(sisters, section.session_time)
        if sister_session:
            overflow += self._write_sisters(sheet, section, sister_session)

        if overflow:
            notes_cell = sheet.cell(row=section.notes_row, column=section.brother_col)
            label = str(notes_cell.value or 'Notes:').rstrip()
            notes_cell.value = f"{label} {', '.join(overflow)}"

    def _session(self, extractor, session_time):
        if extractor is None or session_time is None:
            return None
        return extractor.session_at(session_time)

    def _write_brothers(self, sheet, section, session):
        """Fill the brother column above 'EO:' and both columns below it."""
        names = _display_names(session['presenters'])
        upper_rows = section.upper_rows
        lower_rows = section.lower_rows

        upper_used = min(
            TYPICAL_UPPER_USED if len(names) <= TYPICAL_TOTAL else len(upper_rows),
            len(upper_rows),
        )
        for row, name in zip(upper_rows, names[:upper_used]):
            sheet.cell(row=row, column=section.brother_col).value = name

        remaining = names[upper_used:]
        per_column = min(
            TYPICAL_LOWER_PER_COLUMN
            if len(remaining) <= 2 * TYPICAL_LOWER_PER_COLUMN
            else len(lower_rows),
            len(lower_rows),
        )
        for index, col in enumerate((section.brother_col, section.sister_col)):
            chunk = remaining[index * per_column:(index + 1) * per_column]
            for row, name in zip(lower_rows, chunk):
                sheet.cell(row=row, column=col).value = name

        self._write_officiator(sheet, section, section.brother_col, session['officiator'])
        return remaining[2 * per_column:]

    def _write_sisters(self, sheet, section, session):
        """Fill the sister column above the 'EO:' row."""
        names = _display_names(session['presenters'])
        upper_rows = section.upper_rows

        for row, name in zip(upper_rows, names):
            sheet.cell(row=row, column=section.sister_col).value = name

        self._write_officiator(sheet, section, section.sister_col, session['officiator'])
        return names[len(upper_rows):]

    def _write_officiator(self, sheet, section, col, officiator):
        if not officiator:
            return
        eo_cell = sheet.cell(row=section.eo_row, column=col)
        label = str(eo_cell.value or 'EO:').rstrip()
        eo_cell.value = f"{label} {_display_names([officiator])[0]}"


def _resolve_template(template_path):
    """Locate the template, also checking known project directories."""
    if os.path.exists(template_path):
        return template_path
    if os.path.dirname(template_path):
        return template_path
    for directory in TEMPLATE_SEARCH_DIRS:
        candidate = os.path.join(directory, template_path)
        if os.path.exists(candidate):
            return candidate
    return template_path


def _display_names(names):
    """Shorten 'Last, First' to 'Last', keeping an initial to break ties."""
    last_names = [str(name).split(',', 1)[0].strip() for name in names]
    duplicates = {last for last, count in Counter(last_names).items() if count > 1}

    display = []
    for name, last in zip(names, last_names):
        parts = str(name).split(',', 1)
        first = parts[1].strip() if len(parts) > 1 else ''
        if last in duplicates and first:
            display.append(f"{last}, {first[0]}.")
        else:
            display.append(last)
    return display
