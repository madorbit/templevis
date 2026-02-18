"""
Excel Generator Module

This module handles the generation of Excel spreadsheet files using templates
for temple worker schedule output.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelGenerator:
    """Handles generation of Excel spreadsheet files from schedule data."""

    def __init__(self):
        """Initialize the Excel generator."""
        self.wb = Workbook()
        self.styles = {}
        self._create_styles()

    def _create_styles(self):
        """Create and store common styles for the workbook."""
        # Border styles
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Title style with gray background
        self.styles['title'] = {
            'font': Font(size=14, bold=True),
            'fill': PatternFill(start_color='F0F0F0', end_color='F0F0F0', fill_type='solid'),
            'alignment': Alignment(horizontal='center', vertical='center'),
            'border': thin_border
        }

        # Header style with bold text
        self.styles['header'] = {
            'font': Font(size=14, bold=True),
            'alignment': Alignment(horizontal='left', vertical='center'),
            'border': thin_border
        }

        # Cell style with borders
        self.styles['cell'] = {
            'font': Font(size=14),
            'alignment': Alignment(horizontal='left', vertical='center'),
            'border': thin_border
        }

        # Merged cell style with borders and center alignment
        self.styles['merged_cell'] = {
            'font': Font(size=14),
            'alignment': Alignment(horizontal='center', vertical='center'),
            'border': thin_border
        }

        # Number column style
        self.styles['number_column'] = {
            'font': Font(size=14),
            'alignment': Alignment(horizontal='center', vertical='center'),
            'border': thin_border
        }

        # Alternating row colors
        self.styles['row_white'] = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
        self.styles['row_gray'] = PatternFill(start_color='F0F0F0', end_color='F0F0F0', fill_type='solid')

    def _apply_style(self, cell, style_name):
        """Apply a named style to a cell."""
        style = self.styles[style_name]
        cell.font = style['font']
        if 'fill' in style:
            cell.fill = style['fill']
        cell.alignment = style['alignment']
        cell.border = style['border']

    def create_template_sheet(self, sheet_name, assignment_time):
        """Create a new sheet with the standard template structure."""
        # Remove default sheet if it's the only one
        if len(self.wb.sheetnames) == 1 and self.wb.active.title == "Sheet":
            self.wb.remove(self.wb.active)

        ws = self.wb.create_sheet(sheet_name)

        # Set column widths
        ws.column_dimensions['A'].width = 4   # #
        ws.column_dimensions['B'].width = 25  # Names
        ws.column_dimensions['C'].width = 15  # Room
        ws.column_dimensions['D'].width = 15  # Washing
        ws.column_dimensions['E'].width = 15  # Anointing
        ws.column_dimensions['F'].width = 15  # Clothing

        # Add title row with merged cells (no borders)
        ws.merge_cells('A1:F1')
        title_cell = ws['A1']
        title_cell.value = "Wednesday 4th Shift Initiatory"
        self._apply_style(title_cell, 'title')
        for col in range(1, 7):
            cell = ws.cell(row=1, column=col)
            cell.border = Border(left=Side(style=None), right=Side(style=None), 
                               top=Side(style=None), bottom=Side(style=None))

        # Add empty row (row 2) with smaller height and no borders
        ws.row_dimensions[2].height = 7.5  # 0.1 inches
        for col in range(1, 7):
            cell = ws.cell(row=2, column=col)
            cell.border = Border(left=Side(style=None), right=Side(style=None), 
                               top=Side(style=None), bottom=Side(style=None))

        # Add assignment time row (row 3, no borders)
        ws.merge_cells('A3:B3')
        time_label = ws['A3']
        time_label.value = "Assignment Time:"
        self._apply_style(time_label, 'header')
        time_label.border = Border(left=Side(style=None), right=Side(style=None), 
                                 top=Side(style=None), bottom=Side(style=None))

        ws.merge_cells('C3:F3')
        time_value = ws['C3']
        # Convert 24-hour time to 12-hour format
        try:
            time_parts = assignment_time.split(':')
            hour = int(time_parts[0])
            minute = int(time_parts[1])
            time_str = f"{hour:02d}:{minute:02d} {'PM' if hour >= 12 else 'AM'}"
            if hour > 12:
                time_str = f"{hour-12:02d}:{minute:02d} PM"
            elif hour == 0:
                time_str = f"12:{minute:02d} AM"
            time_value.value = time_str
        except:
            time_value.value = assignment_time
        self._apply_style(time_value, 'header')
        time_value.border = Border(left=Side(style=None), right=Side(style=None), 
                                 top=Side(style=None), bottom=Side(style=None))
        for col in range(3, 7):
            cell = ws.cell(row=3, column=col)
            cell.border = Border(left=Side(style=None), right=Side(style=None), 
                               top=Side(style=None), bottom=Side(style=None))

        # Add empty row (row 4) with smaller height and no borders
        ws.row_dimensions[4].height = 7.5  # 0.1 inches
        for col in range(1, 7):
            cell = ws.cell(row=4, column=col)
            cell.border = Border(left=Side(style=None), right=Side(style=None), 
                               top=Side(style=None), bottom=Side(style=None))

        # Add column headers (row 5)
        headers = ["#", "Names", "Room", "Washing", "Anointing", "Clothing"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col)
            cell.value = header
            self._apply_style(cell, 'header')
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Add numbered rows (6-29) with merged cells for room assignments
        current_room = 1
        rows_per_room = 4
        rows_added = 0

        for i in range(6, 30):  # Rows 6-29 (24 rows)
            # Add row number
            number_cell = ws.cell(row=i, column=1)
            number_cell.value = i - 5  # Numbers 1-24
            self._apply_style(number_cell, 'cell')

            # Add empty name cell
            name_cell = ws.cell(row=i, column=2)
            self._apply_style(name_cell, 'cell')

            # Add room assignment cells and merge cells for other columns
            if rows_added % rows_per_room == 0:
                # First row of a new room group - add merged cells for all columns
                for col in range(3, 7):  # Room, Washing, Anointing, Clothing
                    cell = ws.cell(row=i, column=col)
                    cell.value = current_room if col == 3 else ""  # Only show number in Room column
                    self._apply_style(cell, 'merged_cell')
                    ws.merge_cells(
                        start_row=i,
                        start_column=col,
                        end_row=i + rows_per_room - 1,
                        end_column=col
                    )
                current_room += 1

            rows_added += 1

        return ws

    def add_period_sheet(self, period_name, assignment_time, names, next_tasks=None):
        """
        Add a new sheet for a period with the given names.
        
        Args:
            period_name: Name of the period
            assignment_time: Start time for the period
            names: List of worker names
            next_tasks: Optional dict mapping names to their next task codes
        """
        if not period_name:
            raise ValueError("Sheet name cannot be empty")

        # Convert period name to sheet name format
        sheet_name = period_name.lower()
        if sheet_name == 'period 1':
            sheet_name = '6pm'
        elif sheet_name == 'period 2':
            sheet_name = '645pm'
        elif sheet_name == 'period 3':
            sheet_name = '730pm'
        elif sheet_name == 'period 4':
            sheet_name = '815pm'
        elif sheet_name == 'period 5':
            sheet_name = '9pm'

        ws = self.create_template_sheet(sheet_name, assignment_time)

        # Fill in names starting at row 6
        for i, name in enumerate(names):
            row = i + 6  # Start at row 6
            if row < 30:  # Only fill up to row 29 (24 names)
                # Apply alternating row colors
                row_style = 'row_white' if i % 2 == 0 else 'row_gray'
                for col in range(1, 7):
                    cell = ws.cell(row=row, column=col)
                    cell.fill = self.styles[row_style]

                # Format name with next task decorator if applicable
                display_name = name
                if next_tasks and name in next_tasks:
                    next_task = next_tasks[name]
                    if next_task:
                        display_name = f"{name} (→ {next_task})"

                # Set name and apply style
                name_cell = ws.cell(row=row, column=2)
                name_cell.value = display_name
                self._apply_style(name_cell, 'cell')

                # Apply number column style
                number_cell = ws.cell(row=row, column=1)
                self._apply_style(number_cell, 'number_column')

    def save(self, filepath):
        """Save the workbook to the specified path."""
        if not self.wb:
            raise ValueError("No workbook created yet")
        
        self.wb.save(filepath)
